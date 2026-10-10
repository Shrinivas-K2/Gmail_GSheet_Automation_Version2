"""Real Excel adapter on a temp file; Google adapter against a fake Sheets API."""

import openpyxl

from app.domain.entities import SheetLocation
from app.infrastructure.sheets.excel_gateway import ExcelSheetGateway
from app.infrastructure.sheets import google_gateway as gg
from app.infrastructure.sheets.factory import build_sheet_factory
from app.config.settings import PathSettings, SourceSettings


def make_book(path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Rapido"
    ws.append(["Name", "Email", "Status"])
    ws.append(["Asha", "asha@x.com", None])
    ws.append([None, None, None])
    ws.append(["Bala", "bala@x.com", "Sent"])
    wb.save(path)


def test_excel_gateway_reads_writes_and_creates_columns(tmp_path):
    path = str(tmp_path / "Cold_Mail.xlsx"); make_book(path)
    g = ExcelSheetGateway(path, SheetLocation("excel", "Cold_Mail.xlsx"))
    g.open_tab("Rapido")
    rows = g.read_rows()
    assert [(r.number, r.text("Email")) for r in rows] == [(2, "asha@x.com"), (4, "bala@x.com")]   # blank row skipped
    assert rows[0].get("Status") is None and rows[0].text("Status") == ""

    g.write_cell(2, "Status", "Sent")
    g.write_cell(2, "Reply_Status", "Yes")            # column does not exist -> created
    g.save(); g.close()

    wb = openpyxl.load_workbook(path); ws = wb["Rapido"]
    assert [c.value for c in ws[1]] == ["Name", "Email", "Status", "Reply_Status"]
    assert [c.value for c in ws[2]] == ["Asha", "asha@x.com", "Sent", "Yes"]


def test_excel_gateway_unknown_tab(tmp_path):
    path = str(tmp_path / "b.xlsx"); make_book(path)
    g = ExcelSheetGateway(path, SheetLocation("excel", "b.xlsx"))
    try:
        g.open_tab("Nope"); assert False
    except Exception as e:
        assert str(e) == "Sheet 'Nope' not found."


# ------------------------------------------------------------ Google (fake API)
class X:
    def __init__(self, v): self.v = v
    def execute(self): return self.v


class FakeSheetsApi:
    def __init__(self, tabs): self.tabs, self.writes = tabs, []
    def spreadsheets(self): return self
    def get(self, spreadsheetId, fields): return X({"sheets": [{"properties": {"title": t}} for t in self.tabs]})
    def values(self): return self
    def update(self, spreadsheetId, range, valueInputOption, body):
        self.writes.append((range, body["values"][0][0])); return X({})
    # values().get(...)
    def _values_get(self, spreadsheetId, range): return X({"values": self.tabs[range.strip("'")]})


class ValuesProxy:
    def __init__(self, api): self.api = api
    def get(self, spreadsheetId, range): return self.api._values_get(spreadsheetId, range)
    def update(self, **kw): return self.api.update(**kw)


class SpreadsheetsProxy:
    def __init__(self, api): self.api = api
    def get(self, **kw): return self.api.get(**kw)
    def values(self): return ValuesProxy(self.api)


class ApiClient:
    def __init__(self, api): self.api = api
    def spreadsheets(self): return SpreadsheetsProxy(self.api)


def test_google_gateway_pads_ragged_rows_and_writes_cells():
    api = FakeSheetsApi({"Rapido": [["Name", "Email", "Status"], ["Asha", "asha@x.com"], [], ["Bala", "b@x.com", "Sent"]]})
    g = gg.GoogleSheetGateway(oauth=None, spreadsheet_name="Cold_Mail", spreadsheet_id="ID")
    g._sheets = ApiClient(api)
    g.open_tab("Rapido")
    rows = g.read_rows()
    assert [(r.number, r.get("Status")) for r in rows] == [(2, None), (4, "Sent")]
    g.write_cell(2, "Status", "Sent")
    g.write_cell(2, "Reply_Status", "Yes")
    assert api.writes == [("'Rapido'!C2", "Sent"), ("'Rapido'!D1", "Reply_Status"), ("'Rapido'!D2", "Yes")]
    assert gg._column_letter(27) == "AA" and g.location.spreadsheet_id == "ID"


def test_factory_resolves_default_and_named_locations(tmp_path):
    paths = PathSettings(data_dir=str(tmp_path), excel_file=str(tmp_path / "Cold_Mail.xlsx"))
    f = build_sheet_factory(paths, SourceSettings(google_spreadsheet_name="Sheet1", google_spreadsheet_id="XYZ"), None)
    assert f.default_location().workbook_name == "Cold_Mail.xlsx"
    assert f.default_location("excel", "Other").workbook_name == "Other.xlsx"
    g = f.default_location("google_sheet")
    assert (g.workbook_name, g.spreadsheet_id) == ("Sheet1", "XYZ")
    assert f.default_location("google_sheet", "Named").spreadsheet_id is None
    assert f.create(f.default_location("excel", "Other")).describe() == str(tmp_path / "Other.xlsx")
    try:
        f.default_location("ftp"); assert False
    except ValueError as e:
        assert "Unknown data source" in str(e)
