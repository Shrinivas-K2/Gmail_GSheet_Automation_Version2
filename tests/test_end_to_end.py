"""
Whole program, both command lines, with REAL Excel files and a REAL (SQLite) database.
Only Gmail is faked. This is the safety net for refactoring.
"""

import openpyxl
import pytest

from app.cli import followup_cli, send_cli
from app.config.settings import (
    DatabaseSettings, FollowupSettings, PathSettings, SendingSettings, Settings,
)
from app.container import Container
from app.domain.entities import MailThread, ThreadMessage
from tests.fakes import FakeClock, FakeMailGateway, FakePrompter, FakeReporter

PEOPLE = {
    "Rapido": ["asha", "bala", "chitra", "dev", "esha"],
    "HPE": ["farid"],
}


def make_workbook(path):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for tab, names in PEOPLE.items():
        ws = wb.create_sheet(tab)
        ws.append(["Name", "Email", "Company", "Role", "BCC"])
        for n in names:
            ws.append([n.title(), f"{n}@{tab.lower()}.com", tab, "SDE", "me@gmail.com"])
    wb.save(path)


def read_sheet(path):
    out = {}
    for ws in openpyxl.load_workbook(path):
        headers = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(headers, row)); out[d["Email"]] = d
    return out


@pytest.fixture
def world(tmp_path):
    excel = tmp_path / "Cold_Mail.xlsx"; make_workbook(excel)

    settings = Settings(
        paths=PathSettings(data_dir=str(tmp_path), excel_file=str(excel),
                           log_file=str(tmp_path / "logs" / "mail_log.txt")),
        sending=SendingSettings(attachments=(), delay_seconds=0, retry_pause_seconds=0),
        database=DatabaseSettings(True, f"sqlite:///{tmp_path / 'test.db'}"),
        followup=FollowupSettings(wait_days=(0, 0), check_delay_seconds=0, send_delay_seconds=0),
    )

    def container(followup_gmail=None, prompter=None, sending_gmail=None):
        c = Container(settings, FakeReporter(), FakeClock(), prompter or FakePrompter())
        c.__dict__["sending_gmail"] = sending_gmail or FakeMailGateway()          # override cached_property
        c.__dict__["followup_gmail"] = followup_gmail or FakeMailGateway()
        return c

    return type("World", (), {"excel": str(excel), "settings": settings, "container": staticmethod(container),
                              "tmp": tmp_path})


def me(message_id="<orig@mail>", subject="Subj"):
    return ThreadMessage("me@gmail.com", subject=subject, message_id=message_id, is_outgoing=True)


def build_threads(subject_of):
    """thr1..thr6 = asha, bala, chitra, dev, esha, farid  (order they were sent)."""
    return {
        "thr1": MailThread((me(subject=subject_of), ThreadMessage("Asha <asha@rapido.com>", snippet="Thanks!"))),
        "thr2": MailThread((me(subject=subject_of),)),
        "thr3": MailThread((me(subject=subject_of),)),
        "thr4": MailThread((me(subject=subject_of), ThreadMessage("mailer-daemon@googlemail.com"))),
        "thr5": MailThread((me(subject=subject_of), ThreadMessage("esha@rapido.com", auto_submitted="auto-replied"))),
        "thr6": MailThread((me(subject=subject_of),)),
    }


def send_everything(world):
    sender = FakeMailGateway()
    send_cli.main(["Rapido", "Generic.html"], world.container(sending_gmail=sender))
    send_cli.main(["HPE", "Generic.html"], world.container(sending_gmail=sender))
    return sender


# ======================================================================
def test_send_cli_sends_marks_sheet_and_records_database(world):
    sender = send_everything(world)

    assert [m.to for m in sender.sent] == [f"{n}@{t.lower()}.com" for t, ns in PEOPLE.items() for n in ns]
    sheet = read_sheet(world.excel)
    assert all(row["Status"] == "Sent" for row in sheet.values())

    # a second run sends nothing (already marked Sent)
    again = FakeMailGateway()
    send_cli.main(["Rapido"], world.container(sending_gmail=again))
    assert again.sent == []

    from app.infrastructure.persistence.database import Database
    from app.infrastructure.persistence.orm import EmailRow, RunRow
    db = Database(world.settings.database.url)
    with db.session() as s:
        emails = s.query(EmailRow).all()
        assert len(emails) == 6 and all(e.status == "sent" and e.gmail_thread_id for e in emails)
        assert emails[0].bcc_emails == ["me@gmail.com"] and emails[0].row_snapshot["Name"] == "Asha"
        assert {r.sheet_tab for r in s.query(RunRow)} == {"Rapido", "HPE"}


def test_followup_full_flow_reply_detection_sheet_update_and_threaded_followups(world):
    send_everything(world)
    gmail = FakeMailGateway(threads=build_threads(subject_of="Request"))
    c = world.container(followup_gmail=gmail)
    followup_cli.main(["--all", "--yes"], c, interactive=False)

    sheet = read_sheet(world.excel)
    assert sheet["asha@rapido.com"]["Reply_Status"] == "Yes"
    assert sheet["dev@rapido.com"]["Reply_Status"] == "Bounced"
    assert sheet["bala@rapido.com"]["Reply_Status"] is None
    for who in ("bala@rapido.com", "chitra@rapido.com", "esha@rapido.com", "farid@hpe.com"):
        assert sheet[who]["Followup_Status"].startswith("Follow-up 1 sent")
    assert sheet["asha@rapido.com"]["Followup_Status"] is None and sheet["dev@rapido.com"]["Followup_Status"] is None

    # follow-ups: only the 4 without a real reply / bounce, each INSIDE the original thread
    assert sorted(m.to for m in gmail.sent) == ["bala@rapido.com", "chitra@rapido.com", "esha@rapido.com", "farid@hpe.com"]
    first = next(m for m in gmail.sent if m.to == "bala@rapido.com")
    assert first.thread_id == "thr2" and first.in_reply_to == "<orig@mail>" and first.subject.startswith("Re:")
    assert first.bcc == ["me@gmail.com"] and "Hi <strong>Bala," in first.html_body

    # run 2: follow-up number 2 to the same four.  run 3: the maximum (2) is reached -> nothing
    gmail.sent.clear(); followup_cli.main(["--all", "--yes"], c, interactive=False)
    assert len(gmail.sent) == 4
    gmail.sent.clear(); followup_cli.main(["--all", "--yes"], c, interactive=False)
    assert gmail.sent == []

    from app.infrastructure.persistence.database import Database
    from app.infrastructure.persistence.orm import EmailReplyRow, FollowupRow, FollowupRunRow
    with Database(world.settings.database.url).session() as s:
        assert s.query(FollowupRow).filter_by(status="sent").count() == 8
        counts = {r.email_id: (r.reply_received, r.reply_type, r.followup_count) for r in s.query(EmailReplyRow)}
        assert counts[1] == (True, "reply", 0) and counts[4][1] == "bounce" and counts[2][2] == 2 and counts[5][1] == "auto_reply"
        assert s.query(FollowupRunRow).count() == 3


def test_company_filter_dry_run_and_check_only(world):
    send_everything(world)
    gmail = FakeMailGateway(threads=build_threads("Request"))
    c = world.container(followup_gmail=gmail)

    followup_cli.main(["-c", "hpe", "--dry-run"], c, interactive=False)
    assert gmail.fetched == ["thr6"] and gmail.sent == []
    assert read_sheet(world.excel)["farid@hpe.com"].get("Followup_Status") is None      # nothing written

    gmail.fetched.clear()
    followup_cli.main(["-c", "Rapido", "--check-only"], c, interactive=False)
    assert "thr6" not in gmail.fetched and gmail.sent == []
    assert read_sheet(world.excel)["asha@rapido.com"]["Reply_Status"] == "Yes"        # check-only still updates the sheet

    gmail.fetched.clear()
    followup_cli.main(["-e", "bala@rapido.com", "--limit", "1", "--yes"], c, interactive=False)
    assert gmail.fetched == ["thr2"] and [m.to for m in gmail.sent] == ["bala@rapido.com"]


def test_sending_needs_confirmation_and_respects_no(world):
    send_everything(world)
    gmail = FakeMailGateway(threads=build_threads("Request"))

    # not interactive and no --yes  -> refuses to send
    followup_cli.main(["--all"], world.container(followup_gmail=gmail), interactive=False)
    assert gmail.sent == []

    # interactive, person says no
    no = FakePrompter(confirm=False)
    followup_cli.main(["--all"], world.container(followup_gmail=gmail, prompter=no), interactive=True)
    assert no.asked_confirm == 1 and gmail.sent == []

    # interactive, person picks a company in the question, then says yes
    yes = FakePrompter(scope=["HPE"], confirm=True)
    followup_cli.main([], world.container(followup_gmail=gmail, prompter=yes), interactive=True)
    assert [m.to for m in gmail.sent] == ["farid@hpe.com"]


def test_sync_sheet_repairs_a_sheet_that_missed_updates(world):
    send_everything(world)
    gmail = FakeMailGateway(threads=build_threads("Request"))
    c = world.container(followup_gmail=gmail)
    followup_cli.main(["--all", "--check-only"], c, interactive=False)

    wb = openpyxl.load_workbook(world.excel)                       # wipe Reply_Status, like an Excel that was open
    ws = wb["Rapido"]; col = [x.value for x in ws[1]].index("Reply_Status") + 1
    for r in range(2, ws.max_row + 1): ws.cell(row=r, column=col).value = None
    wb.save(world.excel)

    # a NORMAL run (not --sync-sheet) repairs it, because all known replies are re-written
    followup_cli.main(["--all", "--check-only"], c, interactive=False)
    assert read_sheet(world.excel)["asha@rapido.com"]["Reply_Status"] == "Yes"

    wb = openpyxl.load_workbook(world.excel); ws = wb["Rapido"]
    for r in range(2, ws.max_row + 1): ws.cell(row=r, column=col).value = None
    wb.save(world.excel)
    followup_cli.main(["--sync-sheet", "--all"], c, interactive=False)
    sheet = read_sheet(world.excel)
    assert sheet["asha@rapido.com"]["Reply_Status"] == "Yes" and sheet["dev@rapido.com"]["Reply_Status"] == "Bounced"


def test_unfilled_template_placeholder_blocks_sending_and_is_recorded(world, tmp_path):
    send_everything(world)
    wb = openpyxl.load_workbook(world.excel); ws = wb["HPE"]
    ws.cell(row=2, column=4).value = None           # wipe Role -> Followup_1 needs {Role}
    wb.save(world.excel)
    # the stored snapshot (what was sent) still has Role, so edit the template instead
    tdir = tmp_path / "templates"; tdir.mkdir()
    (tdir / "Followup_1.html").write_text("<!-- SUBJECT: x -->\nHi {Name} {Unknown}")
    from dataclasses import replace
    settings = replace(world.settings, paths=replace(world.settings.paths, templates_dir=str(tdir)))
    c = Container(settings, FakeReporter(), FakeClock(), FakePrompter())
    gmail = FakeMailGateway(threads=build_threads("Request"))
    c.__dict__["followup_gmail"] = gmail
    followup_cli.main(["-c", "HPE", "--yes"], c, interactive=False)
    assert gmail.sent == [] and "Unfilled placeholders" in c.reporter.text

    from app.infrastructure.persistence.database import Database
    from app.infrastructure.persistence.orm import FollowupRow
    with Database(settings.database.url).session() as s:
        row = s.query(FollowupRow).one()
        assert row.status == "failed" and "{Unknown}" in row.error_message


def test_missing_followup_template_stops_early(world):
    send_everything(world)
    gmail = FakeMailGateway(threads=build_threads("Request"))
    c = world.container(followup_gmail=gmail)
    followup_cli.main(["--all", "--yes", "--template", "Nope.html"], c, interactive=False)
    assert "Template not found : templates/Nope.html" in c.reporter.text and gmail.fetched == []


def test_database_down_never_stops_sending_and_followup_explains(world, tmp_path):
    from dataclasses import replace
    bad = replace(world.settings, database=DatabaseSettings(True, "sqlite:////no/such/folder/x.db"))

    c = Container(bad, FakeReporter(), FakeClock(), FakePrompter())
    sender = FakeMailGateway(); c.__dict__["sending_gmail"] = sender
    send_cli.main(["Rapido"], c)
    assert len(sender.sent) == 5 and "[DB] Disabled" in c.reporter.text
    assert all(r["Status"] == "Sent" for e, r in read_sheet(world.excel).items() if e.endswith("rapido.com"))

    c2 = Container(bad, FakeReporter(), FakeClock(), FakePrompter())
    c2.__dict__["followup_gmail"] = FakeMailGateway()
    followup_cli.main(["--all"], c2, interactive=False)
    assert "Cannot connect to the database" in c2.reporter.text


def test_database_off_by_default_uses_null_tracker(world):
    from dataclasses import replace
    off = replace(world.settings, database=DatabaseSettings(False))
    c = Container(off, FakeReporter(), FakeClock(), FakePrompter())
    sender = FakeMailGateway(); c.__dict__["sending_gmail"] = sender
    send_cli.main(["Rapido"], c)
    assert len(sender.sent) == 5 and "Database connected" not in c.reporter.text
