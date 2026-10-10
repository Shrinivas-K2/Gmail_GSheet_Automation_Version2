"""
sheets/excel_gateway.py

SheetGateway for a local .xlsx file (openpyxl).
Changes are kept in memory and written to disk by save() - same as before.
"""

from typing import Dict, List

from openpyxl import load_workbook

from app.domain.entities import SheetLocation, SheetRow
from app.domain.interfaces import SheetGateway


class ExcelSheetGateway(SheetGateway):

    def __init__(self, file_path: str, location: SheetLocation):
        self._file_path = file_path
        self._location = location
        self._workbook = None
        self._sheet = None
        self._header_map: Dict[str, int] = {}

    @property
    def location(self) -> SheetLocation:
        return self._location

    def describe(self) -> str:
        return self._file_path

    # ---- RowReader
    def open_tab(self, tab: str) -> None:

        self._workbook = load_workbook(self._file_path)

        if tab not in self._workbook.sheetnames:
            raise Exception(f"Sheet '{tab}' not found.")

        self._sheet = self._workbook[tab]
        self._load_headers()

    def _load_headers(self) -> None:
        self._header_map = {}

        for column in range(1, self._sheet.max_column + 1):
            value = self._sheet.cell(row=1, column=column).value
            if value is not None:
                self._header_map[str(value).strip()] = column

    def read_rows(self) -> List[SheetRow]:

        rows = []

        for number in range(2, self._sheet.max_row + 1):

            values = {header: self._sheet.cell(row=number, column=column).value
                      for header, column in self._header_map.items()}

            if all(v is None for v in values.values()):
                continue                       # empty row

            rows.append(SheetRow(number, values))

        return rows

    # ---- CellWriter
    def write_cell(self, row_number: int, column: str, value) -> None:

        if column not in self._header_map:
            new_col = self._sheet.max_column + 1
            self._sheet.cell(row=1, column=new_col).value = column
            self._header_map[column] = new_col

        self._sheet.cell(row=row_number, column=self._header_map[column]).value = value

    def save(self) -> None:
        if self._workbook is not None:
            self._workbook.save(self._file_path)

    def close(self) -> None:
        if self._workbook is not None:
            self._workbook.close()
